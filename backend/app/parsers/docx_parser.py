"""
DOCX Parser - Extract text from Word documents using python-docx
"""
from docx import Document
from typing import Optional
import logging
import os

logger = logging.getLogger(__name__)


class DOCXParser:
    """
    DOCX text extraction parser using python-docx.
    Handles various Word document formats.
    """
    
    def __init__(self):
        """Initialize the DOCX parser."""
        pass
    
    def extract_text(self, docx_path: str) -> str:
        """
        Extract all text from a DOCX file.
        
        Args:
            docx_path: Path to the DOCX file
            
        Returns:
            Extracted text as string
            
        Raises:
            FileNotFoundError: If DOCX file doesn't exist
            Exception: For any extraction errors
        """
        text = ""
        
        try:
            # Check if file exists
            if not os.path.exists(docx_path):
                raise FileNotFoundError(f"DOCX file not found: {docx_path}")
            
            doc = Document(docx_path)
            logger.info(f"Processing DOCX file: {docx_path}")
            
            # Extract text from paragraphs
            for para in doc.paragraphs:
                para_text = para.text.strip()
                if para_text:
                    text += para_text + "\n"
            
            # Extract text from tables (common in resumes)
            for table in doc.tables:
                table_text = self._extract_table_text(table)
                if table_text:
                    text += "\n" + table_text + "\n"
            
            # Clean up the extracted text
            text = self._clean_text(text)
            
            logger.info(f"Total extracted text length: {len(text)} characters")
            return text
            
        except FileNotFoundError:
            logger.error(f"DOCX file not found: {docx_path}")
            raise
        except Exception as e:
            logger.error(f"Error extracting text from DOCX: {str(e)}")
            raise
    
    def _extract_table_text(self, table) -> str:
        """
        Extract text from a table, preserving structure.
        
        Args:
            table: python-docx table object
            
        Returns:
            Extracted table text
        """
        table_text = ""
        
        for row in table.rows:
            row_text = []
            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text:
                    row_text.append(cell_text)
            
            if row_text:
                table_text += " | ".join(row_text) + "\n"
        
        return table_text
    
    def _clean_text(self, text: str) -> str:
        """
        Clean and normalize extracted text.
        
        Args:
            text: Raw extracted text
            
        Returns:
            Cleaned text
        """
        if not text:
            return ""
        
        lines = text.split('\n')
        cleaned_lines = []
        
        for line in lines:
            # Strip leading/trailing whitespace
            stripped = line.strip()
            
            # Skip empty lines
            if not stripped:
                continue
            
            # Remove multiple spaces
            cleaned = ' '.join(stripped.split())
            
            if cleaned:
                cleaned_lines.append(cleaned)
        
        # Join lines with proper spacing
        cleaned_text = '\n'.join(cleaned_lines)
        
        return cleaned_text
    
    def extract_metadata(self, docx_path: str) -> dict:
        """
        Extract DOCX metadata.
        
        Args:
            docx_path: Path to the DOCX file
            
        Returns:
            Dictionary containing metadata
        """
        metadata = {}
        
        try:
            doc = Document(docx_path)
            
            # Core properties
            core_props = doc.core_properties
            metadata = {
                'author': core_props.author,
                'title': core_props.title,
                'subject': core_props.subject,
                'keywords': core_props.keywords,
                'last_modified_by': core_props.last_modified_by,
                'revision': core_props.revision,
                'created': str(core_props.created),
                'modified': str(core_props.modified)
            }
            
        except Exception as e:
            logger.error(f"Error extracting metadata: {e}")
            metadata = {'error': str(e)}
        
        return metadata
    
    def extract_headers_footers(self, docx_path: str) -> dict:
        """
        Extract text from headers and footers.
        
        Args:
            docx_path: Path to the DOCX file
            
        Returns:
            Dictionary with header and footer text
        """
        headers_footers = {'headers': [], 'footers': []}
        
        try:
            doc = Document(docx_path)
            
            # Access section headers and footers
            for section in doc.sections:
                # Headers
                header = section.header
                for paragraph in header.paragraphs:
                    if paragraph.text.strip():
                        headers_footers['headers'].append(paragraph.text.strip())
                
                # Footers
                footer = section.footer
                for paragraph in footer.paragraphs:
                    if paragraph.text.strip():
                        headers_footers['footers'].append(paragraph.text.strip())
                        
        except Exception as e:
            logger.error(f"Error extracting headers/footers: {e}")
        
        return headers_footers


# Convenience function for quick extraction
def extract_text_from_docx(docx_path: str) -> str:
    """
    Quick text extraction from DOCX file.
    
    Args:
        docx_path: Path to the DOCX file
        
    Returns:
        Extracted text
    """
    parser = DOCXParser()
    return parser.extract_text(docx_path)
